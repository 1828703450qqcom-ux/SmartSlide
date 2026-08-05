"""
Mock export_service_core module - handles PPTX/PDF export.
"""
import os
import logging
from typing import Optional

logger = logging.getLogger(__name__)


class ExportService:
    """Export service for creating PPTX/PDF files."""

    def __init__(self):
        pass

    def export_to_pptx(self, project_data: dict, output_path: str) -> str:
        """Export project to PPTX format."""
        logger.info("Exporting to PPTX: %s", output_path)
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt

            prs = Presentation()

            pages = project_data.get('pages', [])
            for page in pages:
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)

                title = slide.shapes.title
                if title:
                    title.text = page.get('title', 'Untitled')

                content = slide.placeholders[1]
                if content:
                    content.text = page.get('description', '')

            prs.save(output_path)
            logger.info("PPTX exported successfully: %s", output_path)
            return output_path
        except Exception as e:
            logger.error("Failed to export PPTX: %s", e)
            raise

    def export_to_pdf(self, project_data: dict, output_path: str) -> str:
        """Export project to PDF format."""
        logger.info("Exporting to PDF: %s", output_path)
        raise NotImplementedError("PDF export not yet implemented in mock")


def export_to_pptx(project_data: dict, output_path: str) -> str:
    """Export project to PPTX format (standalone function)."""
    service = ExportService()
    return service.export_to_pptx(project_data, output_path)


def export_to_pdf(project_data: dict, output_path: str) -> str:
    """Export project to PDF format (standalone function)."""
    service = ExportService()
    return service.export_to_pdf(project_data, output_path)
